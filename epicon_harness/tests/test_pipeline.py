"""Offline smoke tests. No network, no SDKs, no keyring.

Run with:  python -m unittest discover -s tests -t . -v
"""

from __future__ import annotations

import json
import tempfile
import unittest
from importlib.util import find_spec
from pathlib import Path

from epicon import paths
from epicon.clients.anthropic_client import AnthropicClient
from epicon.clients.base import ChatResponse, split_thinking
from epicon.conversation import Conversation
from epicon.credentials import StaticCredentialStore
from epicon.errors import ConfigError, SchemaError
from epicon.export_excel import GROUND_TRUTH_FILL, UNDERDETERMINED_FILL, export_sessions
from epicon.gui.render import markdown_to_html
from epicon.model_spec import (
    PINNED_TEMPERATURE,
    STANDARD_MAX_TOKENS,
    THINKING_MAX_TOKENS,
    ModelSpec,
    sanitise_model_id,
)
from epicon.models import (
    STATUS_COMPLETED,
    STATUS_FAILED,
    TURNS_PER_CASE,
    RunLog,
    TestCase,
    Usage,
)
from epicon.runner import BatchRunner, RunnerConfig, fallback_metadata
from epicon.storage import TestCaseIndex, next_run_id, parse_result_path, session_path

TEST_CASE_TEMPLATE = paths.TEMPLATES_DIR / "test_case_template.json"
RUN_LOG_TEMPLATE = paths.TEMPLATES_DIR / "run_log_template.json"
EXAMPLE_CASE = paths.TEST_CASES_DIR / "ALG_021_modified_successor.json"

MODEL = ModelSpec(
    provider="anthropic",
    model_id="Claude-4.6-Opus",
    model_short_id="Claude46",
)

ORACLES = [f"thm_{index:02d}" for index in range(1, TURNS_PER_CASE + 1)]


def make_case(case_id: str = "AA-01") -> TestCase:
    """A sixteen-turn case whose turn N expects the string value of N."""
    return TestCase.from_dict(
        {
            "case_metadata": {
                "case_id": case_id,
                "case_title": f"Title {case_id}",
                "domain": "abstract_algebra",
                "created_by": ["tester"],
                "reviewed_by": [],
                "creation_timestamp_utc": "2026-01-01T00:00:00Z",
            },
            "formal_artifacts": {"lean_file": None, "oracle_theorems": list(ORACLES)},
            "system_instruction": {
                "turn_id": 0,
                "phase": "system_initialization",
                "turn_type": "system_instruction",
                "purpose": "Set the format.",
                "prompt": {"content_raw": "SYSTEM INSTRUCTION", "content_format": "plain"},
                "expected_answer": None,
                "references": None,
            },
            "conversation_framework": [
                {
                    "turn_id": index,
                    "phase": "phase_one" if index < 9 else "phase_two",
                    "turn_type": f"type_{index}",
                    "purpose": f"Purpose {index}.",
                    "prompt": {
                        "content_raw": f"prompt {index}",
                        "content_format": "mixed",
                    },
                    "expected_answer": {
                        "answer_type": "undetermined" if index == 9 else "exact_value",
                        "value": "cannot be determined" if index == 9 else str(index),
                    },
                    "references": None,
                }
                for index in range(1, TURNS_PER_CASE + 1)
            ],
        },
        source=f"<{case_id}>",
    )


class StubClient:
    def __init__(self, *, answers: dict[int, str] | None = None, fail_on_call: int | None = None):
        self.calls: list[list[tuple[str, str]]] = []
        self.systems: list[str] = []
        self.answers = answers or {}
        self.fail_on_call = fail_on_call
        self.closed = False

    def send(self, system, messages, model):
        self.systems.append(system)
        self.calls.append([(message.role, message.content) for message in messages])
        index = len(self.calls)
        if self.fail_on_call is not None and index == self.fail_on_call:
            from epicon.errors import ApiError

            raise ApiError("stub failure", provider=model.provider, attempts=2)
        text = self.answers.get(index, f"Answer: {index}")
        return ChatResponse(
            text=text,
            usage=Usage(input_tokens=10, output_tokens=5, cache_read_input_tokens=3),
            finish_reason="stop",
        )

    def effective_metadata(self, model, *, system_prompt_used: bool):
        return fallback_metadata(model, system_prompt_used=system_prompt_used)

    def close(self) -> None:
        self.closed = True


class StubFactory:
    def __init__(self, **client_kwargs) -> None:
        self.created: list[StubClient] = []
        self.client_kwargs = client_kwargs

    def create(self, provider: str, **_kwargs):
        client = StubClient(**self.client_kwargs)
        self.created.append(client)
        return client


class TestTemplateFidelity(unittest.TestCase):
    def test_test_case_template_round_trips(self) -> None:
        raw = json.loads(TEST_CASE_TEMPLATE.read_text(encoding="utf-8"))
        case = TestCase.from_dict(raw, source=str(TEST_CASE_TEMPLATE))
        self.assertEqual(case.to_dict(), raw)

    def test_run_log_template_round_trips(self) -> None:
        raw = json.loads(RUN_LOG_TEMPLATE.read_text(encoding="utf-8"))
        log = RunLog.from_dict(raw, source=str(RUN_LOG_TEMPLATE))
        self.assertEqual(log.to_dict(), raw)

    def test_test_case_template_semantics(self) -> None:
        case = TestCase.from_file(TEST_CASE_TEMPLATE)
        self.assertEqual(case.case_id, "AA-00")
        self.assertEqual(case.domain, "abstract_algebra")
        self.assertEqual(case.normalised_domain, "abstract-algebra")
        self.assertEqual(case.turn_count, 16)
        self.assertEqual(len(case.formal_artifacts.oracle_theorems), 16)
        self.assertEqual(case.oracle_theorem_for(1), "baseline_successor")
        self.assertTrue(case.system_prompt.startswith("Answer each question"))
        self.assertEqual(case.turn(9).expected_answer.answer_type, "undetermined")
        reference = case.turn(1).references[0]
        self.assertEqual(reference.reference_id, "REF-001")
        self.assertIsNone(case.turn(3).references)

    def test_run_log_template_semantics(self) -> None:
        log = RunLog.from_file(RUN_LOG_TEMPLATE)
        self.assertEqual(log.run_id, "RUN_2026_04_14_0001")
        self.assertEqual(log.run_index, 1)
        self.assertEqual(log.case_id, "AA-01")
        self.assertEqual(log.model_metadata.model_id, "gpt-5.4")
        self.assertEqual(log.model_metadata.temperature, 0.0)
        self.assertIsNone(log.conversation[0].evaluation.is_correct)
        self.assertIsNone(log.conversation[0].evaluation.failure_mode)
        self.assertIsNone(log.conversation[0].evaluation.comments)
        self.assertEqual(log.conversation[0].oracle_theorem_ref, "baseline_successor")
        self.assertIsNone(log.evaluation_summary.correct_turns)

    def test_example_test_case_is_valid(self) -> None:
        case = TestCase.from_file(EXAMPLE_CASE)
        self.assertEqual(case.case_id, "AA-01")
        self.assertEqual(case.turn_count, 16)
        lean_path = case.formal_artifacts.resolved_lean_path()
        self.assertIsNotNone(lean_path)
        self.assertTrue(lean_path.exists())


class TestSchemaValidation(unittest.TestCase):
    def test_missing_case_id_is_rejected(self) -> None:
        with self.assertRaises(SchemaError):
            TestCase.from_dict({"case_metadata": {}, "conversation_framework": []})

    def test_wrong_turn_count_is_rejected(self) -> None:
        with self.assertRaises(SchemaError):
            TestCase.from_dict(
                {
                    "case_metadata": {"case_id": "AA-01", "domain": "abstract_algebra"},
                    "formal_artifacts": {"oracle_theorems": ORACLES},
                    "conversation_framework": [
                        {"turn_id": 1, "turn_type": "t", "prompt": {"content_raw": "x"}}
                    ],
                }
            )

    def test_bad_case_id_is_rejected(self) -> None:
        payload = make_case().to_dict()
        payload["case_metadata"]["case_id"] = "ALG_021"
        with self.assertRaises(SchemaError):
            TestCase.from_dict(payload)

    def test_bad_content_format_is_rejected(self) -> None:
        payload = make_case().to_dict()
        payload["conversation_framework"][0]["prompt"]["content_format"] = "latex"
        with self.assertRaises(SchemaError):
            TestCase.from_dict(payload)


class TestNoRuntimeGrading(unittest.TestCase):
    def test_runner_writes_null_verdicts(self) -> None:
        factory = StubFactory(answers={2: "Answer: 99"})
        with tempfile.TemporaryDirectory() as tmp:
            runner = BatchRunner(
                StaticCredentialStore({"anthropic": "k"}),
                RunnerConfig(model=MODEL, results_dir=Path(tmp) / "results"),
                client_factory=factory,
            )
            log = runner.run_case(make_case())

        for turn in log.conversation:
            self.assertIsNone(turn.evaluation.is_correct)
            self.assertIsNone(turn.evaluation.failure_mode)
            self.assertIsNone(turn.evaluation.comments)
        self.assertEqual(log.conversation[1].evaluation.expected_answer_lean, "2")
        self.assertEqual(log.conversation[0].oracle_theorem_ref, "thm_01")
        self.assertIsNone(log.evaluation_summary.correct_turns)
        self.assertIsNone(log.evaluation_summary.incorrect_turns)

    def test_temperature_is_pinned(self) -> None:
        self.assertEqual(MODEL.temperature, PINNED_TEMPERATURE)
        self.assertEqual(MODEL.token_ceiling, STANDARD_MAX_TOKENS)
        thinking = ModelSpec.from_dict(
            {"provider": "deepseek", "model_id": "DeepSeek-V3.2-Thinking", "thinking": {"enabled": True}}
        )
        self.assertEqual(thinking.token_ceiling, THINKING_MAX_TOKENS)
        self.assertEqual(thinking.thinking_budget, 1024)
        with self.assertRaises(ConfigError):
            ModelSpec.from_dict(
                {"provider": "openai", "model_id": "gpt", "temperature": 0.7}
            )


class TestSanitisedIds(unittest.TestCase):
    def test_spec_examples(self) -> None:
        self.assertEqual(sanitise_model_id("GPT-6-Astra"), "GPT6")
        self.assertEqual(sanitise_model_id("Claude-4.6-Opus"), "Claude46")
        self.assertEqual(sanitise_model_id("DeepSeek-V3.2-Thinking"), "DeepSeekV32")

    def test_session_path_matches_srs(self) -> None:
        path = session_path(
            model_id="Claude-4.6-Opus",
            sanitised_model_id="Claude46",
            run_index=1,
            domain="abstract_algebra",
            case_id="AA-01",
            results_dir=Path("results"),
        )
        self.assertEqual(
            path.as_posix(),
            "results/Claude-4.6-Opus-run-1/abstract-algebra/AA-01-Claude46-R1.json",
        )
        location = parse_result_path(path)
        self.assertIsNotNone(location)
        self.assertEqual(location.model_id, "Claude-4.6-Opus")
        self.assertEqual(location.run_index, 1)
        self.assertEqual(location.domain, "abstract-algebra")
        self.assertEqual(location.case_id, "AA-01")


class TestThinkingSplit(unittest.TestCase):
    def test_think_tags_are_lifted(self) -> None:
        answer, thinking = split_thinking("<think>plan</think>\n4")
        self.assertEqual(answer, "4")
        self.assertEqual(thinking, "plan")

    def test_plain_text_is_untouched(self) -> None:
        answer, thinking = split_thinking("Answer: 4")
        self.assertEqual(answer, "Answer: 4")
        self.assertIsNone(thinking)


class TestConversation(unittest.TestCase):
    def test_turn_order_is_enforced(self) -> None:
        conversation = Conversation(system_prompt="S")
        conversation.append_user("one", 1)
        with self.assertRaises(RuntimeError):
            conversation.append_user("two", 2)
        conversation.append_assistant("reply", 1)
        with self.assertRaises(RuntimeError):
            conversation.messages()
        self.assertEqual(conversation.turn_count, 1)

    def test_rollback_drops_unanswered_prompt(self) -> None:
        conversation = Conversation()
        conversation.append_user("one", 1)
        self.assertIsNotNone(conversation.rollback_pending_user())
        self.assertEqual(len(conversation), 0)


class TestCacheBreakpoints(unittest.TestCase):
    def test_trailing_assistant_turns_are_marked(self) -> None:
        payload = [
            {"role": role, "content": [{"type": "text", "text": "x"}]}
            for role in ("user", "assistant", "user", "assistant", "user")
        ]
        AnthropicClient._apply_cache_breakpoints(payload)
        marked = [
            index
            for index, message in enumerate(payload)
            if "cache_control" in message["content"][-1]
        ]
        self.assertEqual(marked, [1, 3])


class TestRunner(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.results_dir = Path(self._tmp.name) / "results"
        self.journal_dir = Path(self._tmp.name) / "journals"

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def _runner(self, factory: StubFactory, **config_kwargs) -> BatchRunner:
        config = RunnerConfig(
            model=MODEL,
            results_dir=self.results_dir,
            journal_dir=self.journal_dir,
            **config_kwargs,
        )
        return BatchRunner(
            StaticCredentialStore({"anthropic": "test-key"}),
            config,
            client_factory=factory,
        )

    def test_history_is_cumulative(self) -> None:
        factory = StubFactory()
        self._runner(factory).run_case(make_case())
        client = factory.created[0]
        self.assertEqual(len(client.calls), 16)
        self.assertEqual([len(call) for call in client.calls[:3]], [1, 3, 5])
        self.assertEqual(set(client.systems), {"SYSTEM INSTRUCTION"})

    def test_sessions_are_isolated(self) -> None:
        factory = StubFactory()
        result = self._runner(factory).run([make_case("AA-01"), make_case("AA-02")])
        self.assertEqual(len(factory.created), 2)
        self.assertEqual({log.case_id for log in result.logs}, {"AA-01", "AA-02"})

    def test_metadata_is_recorded(self) -> None:
        log = self._runner(StubFactory()).run_case(make_case())
        self.assertEqual(log.status, STATUS_COMPLETED)
        self.assertEqual(log.model_metadata.model_id, "Claude-4.6-Opus")
        self.assertEqual(log.model_metadata.temperature, 0.0)
        self.assertEqual(log.model_metadata.max_tokens, 512)
        self.assertEqual(log.conversation[0].prompt.role, "user")
        self.assertEqual(log.conversation[0].response.role, "assistant")
        self.assertEqual(log.conversation[8].evaluation.expected_answer_type, "undetermined")

    def test_failure_marks_remaining_turns_empty(self) -> None:
        log = self._runner(StubFactory(fail_on_call=2)).run_case(make_case())
        self.assertEqual(log.status, STATUS_FAILED)
        self.assertEqual([turn.turn_id for turn in log.conversation], list(range(1, 17)))
        self.assertEqual(log.conversation[1].response.finish_reason, "error")
        self.assertEqual(log.conversation[2].response.content, "")
        self.assertIsNone(log.conversation[1].evaluation.is_correct)
        self.assertIsNone(log.conversation[2].evaluation.failure_mode)

    def test_missing_model_is_reported(self) -> None:
        runner = BatchRunner(
            StaticCredentialStore({"anthropic": "k"}),
            RunnerConfig(results_dir=self.results_dir),
            client_factory=StubFactory(),
        )
        with self.assertRaises(ConfigError):
            runner.run_case(make_case())

    def test_run_log_lands_in_the_srs_hierarchy(self) -> None:
        self._runner(StubFactory()).run_case(make_case())
        written = list(self.results_dir.rglob("*.json"))
        self.assertEqual(len(written), 1)
        relative = written[0].relative_to(self.results_dir).as_posix()
        self.assertEqual(
            relative, "Claude-4.6-Opus-run-1/abstract-algebra/AA-01-Claude46-R1.json"
        )
        raw = json.loads(written[0].read_text(encoding="utf-8"))
        self.assertEqual(
            list(raw),
            [
                "run_metadata",
                "model_metadata",
                "test_case",
                "conversation",
                "evaluation_summary",
                "researcher_review",
            ],
        )
        reloaded = RunLog.from_file(written[0])
        self.assertEqual(reloaded.to_dict(), raw)
        self.assertEqual(reloaded.usage_totals.output_tokens, 16 * 5)
        self.assertEqual(reloaded.usage_totals.cached_tokens, 16 * 3)

    def test_resume_skips_completed_turns(self) -> None:
        first = self._runner(StubFactory(fail_on_call=8), keep_journal=True)
        failed = first.run_case(make_case())
        self.assertEqual(failed.answered_turns, 7)
        journals = list(self.journal_dir.rglob("*.jsonl"))
        self.assertEqual(len(journals), 1)

        factory = StubFactory()
        resumed = self._runner(factory, resume=True).run_case(make_case())
        self.assertEqual(resumed.status, STATUS_COMPLETED)
        self.assertEqual(resumed.answered_turns, 16)
        # The resumed client only transmitted the remaining nine turns.
        self.assertEqual(len(factory.created[0].calls), 9)
        self.assertEqual(factory.created[0].calls[0][-1], ("user", "prompt 8"))
        self.assertEqual(len(factory.created[0].calls[0]), 15)  # 7 pairs + new prompt

    def test_run_ids_increment(self) -> None:
        runner = self._runner(StubFactory())
        first = runner.run_case(make_case("AA-01")).run_id
        second = runner.run_case(make_case("AA-02")).run_id
        self.assertNotEqual(first, second)

    def test_next_run_id_format(self) -> None:
        self.assertRegex(next_run_id(self.results_dir), r"^RUN_\d{4}_\d{2}_\d{2}_\d{4}$")


class TestRendering(unittest.TestCase):
    def test_display_and_inline_math_survive_markdown(self) -> None:
        html = markdown_to_html("Let $n_1 = n_2$.\n\n$$\\sum_{i=0}^{n} i^2 = x$$\n")
        self.assertIn("n_1 = n_2", html)
        self.assertIn("math-inline", html)
        self.assertIn("math-display", html)

    def test_mid_sentence_display_dollars_render_inline(self) -> None:
        html = markdown_to_html("In this modified system, what is $$2 \\times 3$$?")
        self.assertIn("math-inline", html)
        self.assertNotIn("math-display", html)

    def test_standalone_display_dollars_stay_display(self) -> None:
        html = markdown_to_html("Therefore:\n\n$$3 \\times 4 = 6$$\n")
        self.assertIn("math-display", html)
        self.assertIn("\\[3 \\times 4 = 6\\]", html)

    def test_inspector_splits_prompt_and_completion(self) -> None:
        from epicon.gui.transcript_view import TranscriptRenderer

        log = RunLog.from_file(RUN_LOG_TEMPLATE)
        html = TranscriptRenderer().render(log, None, turn_id=1)
        self.assertIn("inspector", html)
        self.assertIn("Ingested prompt", html)
        self.assertIn("Lean theorem reference", html)
        self.assertIn("Lean ground truth", html)
        self.assertIn("Raw completion", html)
        self.assertIn("thinking", html.lower())

    def test_dollar_inside_code_is_not_math(self) -> None:
        html = markdown_to_html("```bash\necho $HOME and $PATH\n```")
        self.assertNotIn("math-inline", html)
        self.assertIn("$HOME", html)

    @unittest.skipIf(find_spec("markdown") is None, "Markdown not installed")
    def test_markdown_structure_is_rendered(self) -> None:
        html = markdown_to_html("# Heading\n\n- one\n- two\n")
        self.assertIn("<h1>", html)
        self.assertIn("<li>", html)


class TestExport(unittest.TestCase):
    def test_side_by_side_runs_and_highlights(self) -> None:
        try:
            from openpyxl import load_workbook
        except ImportError:  # pragma: no cover
            self.skipTest("openpyxl not installed")

        case = make_case()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            run1 = BatchRunner(
                StaticCredentialStore({"anthropic": "k"}),
                RunnerConfig(model=MODEL, results_dir=root / "r", run_index=1),
                client_factory=StubFactory(),
            ).run_case(case)
            run2 = BatchRunner(
                StaticCredentialStore({"anthropic": "k"}),
                RunnerConfig(model=MODEL, results_dir=root / "r", run_index=2),
                client_factory=StubFactory(answers={1: "other"}),
            ).run_case(case)
            target = root / "out.xlsx"
            result = export_sessions([(run1, case), (run2, case)], target)
            self.assertEqual(result.group_count, 1)
            self.assertEqual(result.turn_count, 16)

            workbook = load_workbook(target)
            self.assertIn("Index", workbook.sheetnames)
            sheet = next(name for name in workbook.sheetnames if name != "Index")
            page = workbook[sheet]
            header = [cell.value for cell in page[1]]
            self.assertIn("Lean Ground Truth", header)
            self.assertIn("Run 1 Output", header)
            self.assertIn("Run 1 Verdict", header)
            self.assertIn("Run 2 Output", header)

            row1 = [cell.value for cell in page[2]]
            self.assertEqual(row1[header.index("Turn")], 1)
            self.assertEqual(row1[header.index("Lean Ground Truth")], "1")
            self.assertEqual(row1[header.index("Run 1 Output")], "Answer: 1")
            self.assertEqual(row1[header.index("Run 2 Output")], "other")
            self.assertIn(row1[header.index("Run 1 Verdict")], ("", None))

            truth = page.cell(row=2, column=header.index("Lean Ground Truth") + 1)
            self.assertEqual((truth.fill.fgColor.rgb or "")[-6:].upper(), GROUND_TRUTH_FILL)

            # Turn 9 is underdetermined.
            yellow = page.cell(row=10, column=1)
            self.assertEqual((yellow.fill.fgColor.rgb or "")[-6:].upper(), UNDERDETERMINED_FILL)

            formula_cell = page.cell(row=2, column=header.index("Run 1 Output") + 1)
            self.assertEqual(formula_cell.data_type, "s")


class TestIndexing(unittest.TestCase):
    def test_logs_pair_with_their_test_case(self) -> None:
        case = make_case("AA-03")
        index = TestCaseIndex([case])
        log = RunLog.for_test_case(case, run_id="RUN_2026_01_01_0001")
        self.assertIs(index.for_log(log), case)
        self.assertEqual(log.test_case.case_title, "Title AA-03")
        self.assertEqual(log.test_case.domain, "abstract_algebra")


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
